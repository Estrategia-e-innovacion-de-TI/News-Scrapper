"""ExcelExporter capability: generates .xlsx reports for ad-hoc queries.

Produces a workbook with two sheets:
- "Resultados": one row per document with classification data
- "Metadata": run-level metadata (run_id, dates, terms, counts)

Registered in CapabilityRegistry as ``export_excel``.
"""
from __future__ import annotations

import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from extractor.capabilities.registry import CapabilityDef, register_capability
from extractor.state import Document

logger = logging.getLogger(__name__)

# Column definitions for "Resultados" sheet
RESULTADOS_COLUMNS = [
    "título",
    "medio",
    "fecha",
    "URL",
    "resumen_corto",
    "categoría",
    "severidad",
    "evidencia_citas",
]

# Field definitions for "Metadata" sheet
METADATA_FIELDS = [
    "run_id",
    "fecha_ejecución",
    "empresa_o_términos",
    "rango_fechas",
    "total_documentos",
    "total_clasificados",
]

MAX_COLUMN_WIDTH = 80
MAX_EVIDENCE_SPANS = 3
EVIDENCE_SEPARATOR = " | "


def _format_evidence(doc: Document) -> str:
    """Format evidence_spans as pipe-separated text, max 3 spans."""
    spans = doc.evidence_spans[:MAX_EVIDENCE_SPANS]
    if not spans:
        return ""
    return EVIDENCE_SEPARATOR.join(span.text for span in spans)


def _auto_adjust_column_widths(ws: Any) -> None:
    """Auto-adjust column widths based on content, capped at MAX_COLUMN_WIDTH."""
    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1), start=1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_len:
                    max_len = cell_len
        # Add a small padding, cap at MAX_COLUMN_WIDTH
        adjusted = min(max_len + 2, MAX_COLUMN_WIDTH)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted


def _build_resultados_sheet(ws: Any, documents: list[Document]) -> None:
    """Populate the 'Resultados' sheet with document rows."""
    # Write headers
    ws.append(RESULTADOS_COLUMNS)

    if not documents:
        ws.append(["Sin resultados para los criterios especificados"])
        return

    for doc in documents:
        ws.append([
            doc.title,
            doc.source_id,
            doc.published_at or "",
            doc.url,
            doc.excerpt[:200] if doc.excerpt else "",
            doc.category or doc.risk_type or "",
            doc.severity or "",
            _format_evidence(doc),
        ])


def _build_metadata_sheet(ws: Any, metadata: dict[str, Any]) -> None:
    """Populate the 'Metadata' sheet with run-level info."""
    ws.append(METADATA_FIELDS)
    ws.append([
        metadata.get("run_id", ""),
        metadata.get("fecha_ejecución", datetime.now().isoformat()),
        metadata.get("empresa_o_términos", ""),
        metadata.get("rango_fechas", ""),
        metadata.get("total_documentos", 0),
        metadata.get("total_clasificados", 0),
    ])


def export_excel(
    documents: list[Document],
    metadata: dict[str, Any],
    output_path: str | Path,
) -> Path:
    """Generate an Excel report from classified documents.

    Parameters
    ----------
    documents:
        List of Document objects (may be empty).
    metadata:
        Dict with keys: run_id, fecha_ejecución, empresa_o_términos,
        rango_fechas, total_documentos, total_clasificados.
    output_path:
        File path for the .xlsx output (directories created if needed).

    Returns
    -------
    Path to the written .xlsx file.
    """
    output_path = Path(output_path)

    # Create output directory if it doesn't exist
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # --- Sheet 1: Resultados ---
    ws_resultados = wb.active
    ws_resultados.title = "Resultados"
    _build_resultados_sheet(ws_resultados, documents)
    _auto_adjust_column_widths(ws_resultados)

    # --- Sheet 2: Metadata ---
    ws_metadata = wb.create_sheet(title="Metadata")
    _build_metadata_sheet(ws_metadata, metadata)
    _auto_adjust_column_widths(ws_metadata)

    wb.save(str(output_path))
    logger.info("Excel report written to %s (%d documents)", output_path, len(documents))
    return output_path


# ---------------------------------------------------------------------------
# Register in CapabilityRegistry
# ---------------------------------------------------------------------------
register_capability(
    CapabilityDef(
        name="export_excel",
        purpose="Generate .xlsx report for ad-hoc ARAS/Riesgos queries",
        inputs_schema={
            "documents": "list[Document]",
            "metadata": "dict[str, Any]",
            "output_path": "str | Path",
        },
        outputs_schema={"path": "Path"},
        callable=export_excel,
        version="1.0",
        tags=["export", "excel", "adhoc"],
    )
)
