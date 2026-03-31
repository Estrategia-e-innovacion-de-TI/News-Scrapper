"""ExcelExporter — generates .xlsx reports for ad-hoc queries.

Produces a workbook with two sheets:
- "Resultados": one row per document with classification data
- "Metadata": run-level metadata (run_id, dates, terms, counts)

Ported from news_radar_mvp/extractor/capabilities/excel_export.py
Validates: Requirements 12.1-12.6
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from newsradar_api.domain.model.pipeline_models import Document

logger = logging.getLogger(__name__)

# Column definitions for "Resultados" sheet
RESULTADOS_COLUMNS = [
    "título",
    "medio",
    "fecha",
    "URL",
    "resumen_corto",
    "categoría",
    "severidad",
    "evidencia_citas",
]

# Field definitions for "Metadata" sheet
METADATA_FIELDS = [
    "run_id",
    "fecha_ejecución",
    "empresa_o_términos",
    "rango_fechas",
    "total_documentos",
    "total_clasificados",
]

MAX_COLUMN_WIDTH = 80
MAX_EVIDENCE_SPANS = 3
EVIDENCE_SEPARATOR = " | "


def _format_evidence(doc: Document) -> str:
    """Format evidence_spans as pipe-separated text, max 3 spans.

    Validates: Requirement 12.2 (evidencia_citas máx 3 spans separated by " | ")
    """
    spans = doc.evidence_spans[:MAX_EVIDENCE_SPANS]
    if not spans:
        return ""
    return EVIDENCE_SEPARATOR.join(span.text for span in spans)


def _auto_adjust_column_widths(ws: Any) -> None:
    """Auto-adjust column widths based on content, capped at MAX_COLUMN_WIDTH.

    Validates: Requirement 12.4
    """
    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1), start=1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_len:
                    max_len = cell_len
        adjusted = min(max_len + 2, MAX_COLUMN_WIDTH)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted


def _build_resultados_sheet(ws: Any, documents: list[Document]) -> None:
    """Populate the 'Resultados' sheet with document rows.

    Validates: Requirements 12.2, 12.6
    """
    ws.append(RESULTADOS_COLUMNS)

    # Req 12.6: empty list → single row with message
    if not documents:
        ws.append(["Sin resultados para los criterios especificados"])
        return

    for doc in documents:
        ws.append([
            doc.title,
            doc.source_id,
            doc.published_at or "",
            doc.url,
            doc.excerpt[:200] if doc.excerpt else "",
            doc.category or doc.risk_type or "",
            doc.severity or "",
            _format_evidence(doc),
        ])


def _build_metadata_sheet(ws: Any, metadata: dict[str, Any]) -> None:
    """Populate the 'Metadata' sheet with run-level info.

    Validates: Requirement 12.3
    """
    ws.append(METADATA_FIELDS)
    ws.append([
        metadata.get("run_id", ""),
        metadata.get("fecha_ejecución", datetime.now().isoformat()),
        metadata.get("empresa_o_términos", ""),
        metadata.get("rango_fechas", ""),
        metadata.get("total_documentos", 0),
        metadata.get("total_clasificados", 0),
    ])


class ExcelExporter:
    """Generates .xlsx reports from classified documents.

    Validates: Requirements 12.1-12.6
    """

    def export(
        self,
        docs: list[Document],
        output_path: str | Path,
        metadata: dict[str, Any] | None = None,
    ) -> Path:
        """Generate an Excel report from classified documents.

        Parameters
        ----------
        docs
            List of Document objects (may be empty).
        output_path
            File path for the .xlsx output (directories created if needed).
        metadata
            Dict with keys: run_id, fecha_ejecución, empresa_o_términos,
            rango_fechas, total_documentos, total_clasificados.

        Returns
        -------
        Path to the written .xlsx file.
        """
        output_path = Path(output_path)
        if metadata is None:
            metadata = {}

        # Req 12.5: create parent directories if they don't exist
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()

        # Req 12.1: Sheet "Resultados"
        ws_resultados = wb.active
        ws_resultados.title = "Resultados"
        _build_resultados_sheet(ws_resultados, docs)
        _auto_adjust_column_widths(ws_resultados)

        # Req 12.1: Sheet "Metadata"
        ws_metadata = wb.create_sheet(title="Metadata")
        _build_metadata_sheet(ws_metadata, metadata)
        _auto_adjust_column_widths(ws_metadata)

        wb.save(str(output_path))
        logger.info(
            "Excel report written to %s (%d documents)", output_path, len(docs),
        )
        return output_path


__all__ = ["ExcelExporter"]
