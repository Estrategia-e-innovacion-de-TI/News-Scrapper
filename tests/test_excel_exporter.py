"""Unit tests for ExcelExporter capability."""
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from extractor.capabilities.excel_export import (
    EVIDENCE_SEPARATOR,
    MAX_COLUMN_WIDTH,
    METADATA_FIELDS,
    RESULTADOS_COLUMNS,
    export_excel,
    _format_evidence,
)
from extractor.capabilities.registry import get_capability
from extractor.state import Document, EvidenceSpan


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_doc(
    *,
    title: str = "Noticia de prueba",
    source_id: str = "src_1",
    url: str = "https://example.com/1",
    published_at: str | None = "2025-01-15",
    excerpt: str = "Resumen corto de la noticia",
    category: str | None = "fraude",
    severity: str | None = "H",
    evidence_spans: list[EvidenceSpan] | None = None,
    text: str = "Texto completo de la noticia de prueba.",
) -> Document:
    return Document(
        run_id="test-run",
        source_id=source_id,
        pipeline_class="news",
        focus=["aras_news"],
        title=title,
        url=url,
        canonical_url=None,
        published_at=published_at,
        fetched_at=datetime.now().isoformat(),
        language="es",
        text=text,
        excerpt=excerpt,
        raw_len=len(text),
        text_len=len(text),
        hash="abc123",
        fetch_method="http",
        source_url="https://example.com",
        category=category,
        severity=severity,
        evidence_spans=evidence_spans or [],
    )


def _make_metadata(**overrides: object) -> dict:
    base = {
        "run_id": "test-run",
        "fecha_ejecución": "2025-01-15T10:00:00",
        "empresa_o_términos": "Empresa Test",
        "rango_fechas": "2025-01-01 / 2025-01-15",
        "total_documentos": 5,
        "total_clasificados": 4,
    }
    base.update(overrides)
    return base


# ---------------------------------------------------------------------------
# Registration
# ---------------------------------------------------------------------------

class TestCapabilityRegistration:
    def test_export_excel_registered(self):
        cap = get_capability("export_excel")
        assert cap is not None
        assert cap.name == "export_excel"
        assert cap.callable.__name__ == export_excel.__name__
        assert cap.callable.__module__.endswith("excel_export")


# ---------------------------------------------------------------------------
# Sheet structure
# ---------------------------------------------------------------------------

class TestResultadosSheet:
    def test_headers_present(self, tmp_path: Path):
        path = tmp_path / "out" / "results.xlsx"
        export_excel([], _make_metadata(), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        headers = [cell.value for cell in ws[1]]
        assert headers == RESULTADOS_COLUMNS

    def test_document_row_values(self, tmp_path: Path):
        spans = [
            EvidenceSpan(text="...evidencia [fraude]...", start_offset=10, end_offset=30, matched_term="fraude"),
        ]
        doc = _make_doc(evidence_spans=spans)
        path = tmp_path / "results.xlsx"
        export_excel([doc], _make_metadata(), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        row = [cell.value for cell in ws[2]]
        assert row[0] == doc.title
        assert row[1] == doc.source_id
        assert row[2] == doc.published_at
        assert row[3] == doc.url
        assert row[5] == "fraude"
        assert row[6] == "H"
        assert "evidencia" in row[7]

    def test_multiple_documents(self, tmp_path: Path):
        docs = [_make_doc(title=f"Doc {i}") for i in range(3)]
        path = tmp_path / "results.xlsx"
        export_excel(docs, _make_metadata(total_documentos=3), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        # 1 header + 3 data rows
        assert ws.max_row == 4


class TestMetadataSheet:
    def test_metadata_headers(self, tmp_path: Path):
        path = tmp_path / "results.xlsx"
        export_excel([], _make_metadata(), path)
        wb = load_workbook(path)
        ws = wb["Metadata"]
        headers = [cell.value for cell in ws[1]]
        assert headers == METADATA_FIELDS

    def test_metadata_values(self, tmp_path: Path):
        meta = _make_metadata()
        path = tmp_path / "results.xlsx"
        export_excel([], meta, path)
        wb = load_workbook(path)
        ws = wb["Metadata"]
        row = [cell.value for cell in ws[2]]
        assert row[0] == "test-run"
        assert row[2] == "Empresa Test"
        assert row[4] == 5
        assert row[5] == 4


# ---------------------------------------------------------------------------
# Evidence formatting
# ---------------------------------------------------------------------------

class TestEvidenceFormatting:
    def test_no_spans(self):
        doc = _make_doc(evidence_spans=[])
        assert _format_evidence(doc) == ""

    def test_single_span(self):
        spans = [EvidenceSpan(text="span1", start_offset=0, end_offset=5, matched_term="t")]
        doc = _make_doc(evidence_spans=spans)
        assert _format_evidence(doc) == "span1"

    def test_three_spans_pipe_separated(self):
        spans = [
            EvidenceSpan(text=f"span{i}", start_offset=i * 10, end_offset=i * 10 + 5, matched_term="t")
            for i in range(3)
        ]
        doc = _make_doc(evidence_spans=spans)
        result = _format_evidence(doc)
        assert result == "span0" + EVIDENCE_SEPARATOR + "span1" + EVIDENCE_SEPARATOR + "span2"

    def test_max_three_spans_truncated(self):
        spans = [
            EvidenceSpan(text=f"span{i}", start_offset=i * 10, end_offset=i * 10 + 5, matched_term="t")
            for i in range(5)
        ]
        doc = _make_doc(evidence_spans=spans)
        result = _format_evidence(doc)
        parts = result.split(EVIDENCE_SEPARATOR)
        assert len(parts) == 3


# ---------------------------------------------------------------------------
# Empty results
# ---------------------------------------------------------------------------

class TestEmptyResults:
    def test_empty_list_produces_note(self, tmp_path: Path):
        path = tmp_path / "results.xlsx"
        export_excel([], _make_metadata(total_documentos=0), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        # Row 1 = headers, Row 2 = note
        assert ws.max_row == 2
        note = ws.cell(row=2, column=1).value
        assert note == "Sin resultados para los criterios especificados"


# ---------------------------------------------------------------------------
# Output directory creation
# ---------------------------------------------------------------------------

class TestDirectoryCreation:
    def test_creates_nested_dirs(self, tmp_path: Path):
        path = tmp_path / "a" / "b" / "c" / "results.xlsx"
        assert not path.parent.exists()
        export_excel([], _make_metadata(), path)
        assert path.exists()


# ---------------------------------------------------------------------------
# Column widths
# ---------------------------------------------------------------------------

class TestColumnWidths:
    def test_widths_capped_at_max(self, tmp_path: Path):
        long_title = "A" * 200
        doc = _make_doc(title=long_title)
        path = tmp_path / "results.xlsx"
        export_excel([doc], _make_metadata(), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        for col_idx in range(1, ws.max_column + 1):
            letter = chr(64 + col_idx) if col_idx <= 26 else None
            if letter:
                width = ws.column_dimensions[letter].width
                assert width <= MAX_COLUMN_WIDTH


# ---------------------------------------------------------------------------
# Return value
# ---------------------------------------------------------------------------

class TestReturnValue:
    def test_returns_path(self, tmp_path: Path):
        path = tmp_path / "results.xlsx"
        result = export_excel([], _make_metadata(), path)
        assert result == path
        assert isinstance(result, Path)


# ---------------------------------------------------------------------------
# Category / risk_type fallback
# ---------------------------------------------------------------------------

class TestCategoryFallback:
    def test_risk_type_used_when_no_category(self, tmp_path: Path):
        doc = _make_doc(category=None)
        doc.risk_type = "cibernetico"
        path = tmp_path / "results.xlsx"
        export_excel([doc], _make_metadata(), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        cat_cell = ws.cell(row=2, column=6).value
        assert cat_cell == "cibernetico"

    def test_both_none_gives_empty(self, tmp_path: Path):
        doc = _make_doc(category=None, severity=None)
        path = tmp_path / "results.xlsx"
        export_excel([doc], _make_metadata(), path)
        wb = load_workbook(path)
        ws = wb["Resultados"]
        cat_cell = ws.cell(row=2, column=6).value
        assert cat_cell in ("", None)
